from pathlib import Path
import pandas as pd
from openpyxl.styles import Font, PatternFill, Alignment

root = Path(__file__).resolve().parents[1]
out = root / 'Datos_CAMELS.xlsx'
r = root / 'resultados'
all_data = pd.read_csv(r / 'revision_todas_cuencas.csv')
notes = [
    ('Contenido', 'Datos CAMELS-COL y revisión preliminar para seleccionar una cuenca.'),
    ('Fuente original', 'https://doi.org/10.5281/zenodo.18794895'),
    ('Versión del depósito', '26 de febrero de 2026'),
    ('Fecha de consulta', '21 de septiembre de 2026'),
    ('Identificación de estaciones', 'Nombres y ubicación complementados con el catálogo IDEAM: https://www.datos.gov.co/resource/hp9r-jxuu.json'),
    ('Filtro', 'Área de 100 a 10000 km² y máximo 10 % de días faltantes por variable P y Q en cada periodo evaluado.'),
    ('Resultado', '102 cuencas cumplen el filtro en 1981–2022. La elección definitiva sigue pendiente.'),
    ('Precipitación diaria', 'CHIRPS v2, mm/día; no es una medición puntual de pluviómetro.'),
    ('Caudal diario', 'Caudal observado reportado por CAMELS, m³/s. No hay banderas de calidad por observación en estos archivos.'),
    ('Temperaturas', 'Mínima y máxima de MSWX, °C. No se encontró una columna de temperatura media.'),
    ('ETP_', 'Evapotranspiración potencial diaria reportada, mm/día; se conserva sin validación específica.'),
    ('Fechas ausentes', 'Las hojas diarias conservan las filas originales; los días omitidos se contabilizan al reconstruir el calendario en la revisión.'),
    ('Faltantes', 'Calculados antes de cualquier relleno sobre todos los días esperados, incluidos años bisiestos.'),
    ('Mes completo', '100 % de días válidos; las series mensuales preliminares dejan vacíos los meses incompletos.'),
    ('Coordenadas', 'gauge_lat/gauge_lon originales están en metros EPSG:3395; latitud/longitud del catálogo IDEAM están en grados.'),
    ('IMERG', 'Disponibilidad en catálogo verificada; no se incluyen valores IMERG descargados para las cuencas.'),
    ('Limitación', 'Cumplir el filtro no certifica ausencia de errores, rellenos previos o regulación. Consultar README.md.'),
]
with pd.ExcelWriter(out, engine='openpyxl') as writer:
    pd.DataFrame(notes, columns=['Tema', 'Descripción']).to_excel(writer, sheet_name='Leer primero', index=False)
    all_data[(all_data.periodo == '1981_2022') & all_data.gauge_id.isin([12027050,22027020,26127040])].to_excel(writer,sheet_name='Tres candidatas',index=False)
    pd.read_csv(r/'candidatas_42_anos.csv').to_excel(writer,sheet_name='102 cuencas aptas',index=False)
    all_data.to_excel(writer,sheet_name='Revisión 346 cuencas',index=False)
    pd.read_csv(r/'disponibilidad_mensual_candidatas.csv').to_excel(writer,sheet_name='Disponibilidad mensual',index=False)
    for code in [12027050,22027020,26127040,23057140,21017020]:
        for kind,label in [('diario','Diario'),('mensual_preliminar','Mensual')]:
            frame=pd.read_csv(r/f'{kind}_{code}.csv')
            frame[frame.columns[0]]=pd.to_datetime(frame.iloc[:,0])
            frame.to_excel(writer,sheet_name=f'{label} {code}',index=False)
    for sheet in writer.book:
        sheet.freeze_panes='A2'
        sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:
            cell.font=Font(color='FFFFFF',bold=True)
            cell.fill=PatternFill('solid',fgColor='176B87')
        for column in sheet.columns:
            letter=column[0].column_letter
            sheet.column_dimensions[letter].width=min(38,max(16,len(str(column[0].value))+2))
    intro=writer.book['Leer primero']
    intro.column_dimensions['A'].width=27
    intro.column_dimensions['B'].width=110
    for row in intro.iter_rows(min_row=2):
        row[1].alignment=Alignment(wrap_text=True,vertical='top')
        intro.row_dimensions[row[0].row].height=34
print(out)
