"""Cruce espacial reproducible del catalogo IDEAM con ocho cuencas CAMELS.
Los conteos son de estaciones, no de series disponibles ni de sensores verificados.
"""
from pathlib import Path
import json
import shutil
import unicodedata
import pandas as pd
from shapely.geometry import shape, Point
from openpyxl import load_workbook
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.styles import Font, PatternFill, Alignment

ROOT=Path(__file__).resolve().parents[1]
D,R=ROOT/'datos',ROOT/'resultados'
def normal(text):
    return ''.join(c for c in unicodedata.normalize('NFKD',str(text)) if not unicodedata.combining(c)).lower().strip()

catalog=pd.DataFrame(json.loads((D/'catalogo_ideam.json').read_text(encoding='utf-8')))
if catalog.codigo.duplicated().any():
    raise ValueError('Catalogo con codigos duplicados: revisar antes del conteo')
# Usar columnas explicitas: ubicaci_n tiene lat/lon intercambiadas en algunos registros.
catalog['latitud_num']=pd.to_numeric(catalog.latitud,errors='coerce')
catalog['longitud_num']=pd.to_numeric(catalog.longitud,errors='coerce')
valid=catalog.latitud_num.between(-90,90)&catalog.longitud_num.between(-180,180)
catalog.loc[~valid].to_csv(R/'estaciones_coordenadas_invalidas.csv',index=False,encoding='utf-8-sig')
catalog=catalog.loc[valid].copy()
catalog['cat_norm']=catalog.categoria.map(normal)
catalog['activa']=catalog.estado.eq('Activa')
rain={'pluviometrica','pluviografica','climatologica ordinaria','climatologica principal','agrometeorologica','sinoptica principal','sinoptica secundaria'}
meteo_multi={'climatologica ordinaria','climatologica principal','agrometeorologica','sinoptica principal','sinoptica secundaria','meteorologica especial','meteorologica marina','hidrometeorologica'}
catalog['pluvial_exclusiva']=catalog.cat_norm.isin({'pluviometrica','pluviografica'})
catalog['meteorologica_multivariable']=catalog.cat_norm.isin(meteo_multi)
catalog['lluvia_por_categoria']=catalog.cat_norm.isin(rain)
catalog['hidrometrica']=catalog.cat_norm.isin({'limnimetrica','limnigrafica'})
geometries=json.loads((R/'ocho_cuencas_wgs84.geojson').read_text(encoding='utf-8'))['features']
summary=pd.read_csv(R/'filtro_8_temperatura_imerg.csv')
audits=pd.read_csv(R/'revision_todas_cuencas.csv')
details=[]
for feature in geometries:
    code=int(feature['properties']['gauge_id'])
    poly=shape(feature['geometry'])
    w,s,e,n=poly.bounds
    near=catalog[catalog.longitud_num.between(w,e)&catalog.latitud_num.between(s,n)]
    inside=near[[poly.covers(Point(lon,lat)) for lon,lat in zip(near.longitud_num,near.latitud_num)]].copy()
    inside.insert(0,'cuenca_codigo',code)
    inside['en_borde']= [poly.boundary.distance(Point(lon,lat))<1e-9 for lon,lat in zip(inside.longitud_num,inside.latitud_num)]
    details.append(inside)
    idx=summary.index[summary.codigo.eq(code)][0]
    for key,flag in [('pluviometricas_pluviograficas','pluvial_exclusiva'),('meteorologicas_multivariable','meteorologica_multivariable'),('lluvia_segun_categoria','lluvia_por_categoria'),('hidrometricas','hidrometrica')]:
        summary.loc[idx,key+'_total']=int(inside[flag].sum())
        summary.loc[idx,key+'_activas']=int((inside[flag]&inside.activa).sum())
    summary.loc[idx,'todas_estaciones_total']=len(inside)
    for period in ['1981_2022','1998_2022']:
        a=audits[(audits.gauge_id==code)&(audits.periodo==period)].iloc[0]
        for var in ['P','Q','Tmin','Tmax']:
            summary.loc[idx,f'faltantes_{var}_pct_{period}']=a[f'faltantes_{var}_pct']
    summary.loc[idx,'faltante_max_PQT_1998_2022_pct']=max(summary.loc[idx,f'faltantes_{v}_pct_1998_2022'] for v in ['P','Q','Tmin','Tmax'])

for col in summary:
    if col.endswith(('_total','_activas')):
        summary[col]=summary[col].astype(int)
summary=summary.sort_values(['lluvia_segun_categoria_activas','faltante_max_PQT_1998_2022_pct','lluvia_segun_categoria_total','codigo'],ascending=[False,True,False,True]).reset_index(drop=True)
summary.insert(0,'prioridad',range(1,len(summary)+1))
summary.insert(1,'candidata_principal',summary.index==0)
# Frontera de Pareto: explicita el compromiso entre mas estaciones y menos vacios.
summary['pareto_estaciones_faltantes']=[not any((other.lluvia_segun_categoria_activas>=row.lluvia_segun_categoria_activas and other.faltante_max_PQT_1998_2022_pct<=row.faltante_max_PQT_1998_2022_pct and (other.lluvia_segun_categoria_activas>row.lluvia_segun_categoria_activas or other.faltante_max_PQT_1998_2022_pct<row.faltante_max_PQT_1998_2022_pct)) for _,other in summary.iterrows()) for _,row in summary.iterrows()]
detail=pd.concat(details,ignore_index=True)
cols=['cuenca_codigo','codigo','nombre','categoria','estado','tecnologia','departamento','municipio','latitud_num','longitud_num','fecha_instalacion','activa','pluvial_exclusiva','meteorologica_multivariable','lluvia_por_categoria','hidrometrica','en_borde']
detail=detail[cols].sort_values(['cuenca_codigo','categoria','nombre'])
summary.to_csv(R/'ranking_estaciones_ideam.csv',index=False,encoding='utf-8-sig')
detail.to_csv(R/'estaciones_dentro_cuencas.csv',index=False,encoding='utf-8-sig')
notes=pd.DataFrame([
    ('Fuente','Catalogo IDEAM descargado el 21/09/2026: https://www.datos.gov.co/resource/hp9r-jxuu.json; '+str(len(catalog))+' estaciones con coordenadas validas.'),
    ('Cruce espacial','Punto dentro o en borde del poligono CAMELS, transformado a WGS84. Sin buffer. Coordenadas explicitas latitud/longitud; no se usa ubicaci_n.'),
    ('Meteorologicas multivariable','Climatologicas ordinarias/principales, agrometeorologicas, sinopticas, meteorologicas especiales/marinas e hidrometeorologicas. Separadas de PM/PG para evitar sumarlas dos veces.'),
    ('Lluvia segun categoria','PM + PG + climatologicas ordinarias/principales + agrometeorologicas + sinopticas. Se infiere capacidad por categoria, no se verificaron sensores ni series individuales.'),
    ('Fuentes de clasificacion','https://www.ideam.gov.co/sites/default/files/mapa-de-procesos/gdi-g005_guia_para_la_operacion_y_mantenimiento_de_las_estaciones_meteorologicas_convencionales_v3.pdf ; https://bart.ideam.gov.co/meteorologia/NOTAS_TECNICAS/NT_IDEAM-2006-004_LA_SEQU%C3%8DA_EN_COLOMBIA.pdf'),
    ('Totales y activas','Total incluye activas, suspendidas y en mantenimiento. Activas usa exclusivamente el estado Activa del catalogo guardado; no garantiza que estuvieran operando durante todo 1998-2022.'),
    ('Regla de prioridad','Primero mayor numero de estaciones activas de lluvia segun categoria; segundo menor maximo de faltantes P/Q/Tmin/Tmax CAMELS en 1998-2022; despues mayor total historico de estaciones de lluvia.'),
    ('Compromiso','Mas estaciones y menos faltantes pueden favorecer cuencas diferentes. Se incluye frontera de Pareto; no se afirma que la principal minimice tambien los faltantes.'),
    ('Faltantes','Son faltantes de las series CAMELS de la cuenca, NO de las estaciones meteorologicas del inventario. Faltantes de esas estaciones: no evaluados.'),
    ('Periodo','1998-2022 = 25 anos comunes para evaluar disponibilidad CAMELS y cobertura catalogada IMERG. Tambien se incluyen faltantes CAMELS 1981-2022.'),
    ('Sin duplicados','Cada codigo se cuenta una sola vez por cuenca. Lluvia segun categoria se solapa con meteorologicas multivariable: NO sumar ambas columnas.'),
    ('Cobertura','Solo catalogo IDEAM; pueden existir estaciones de otras redes no representadas. Mayor numero no garantiza mejor distribucion espacial ni representatividad.'),
    ('IMERG','Continua disponible solo en catalogo; no se han descargado valores. Su disponibilidad se conserva en las hojas originales.'),
],columns=['Tema','Explicacion'])
book=ROOT/'Filtro_8_cuencas_IMERG.xlsx'
backup=ROOT/'Filtro_8_cuencas_IMERG_antes_estaciones.xlsx'
if not backup.exists():
    shutil.copy2(book,backup)
wb=load_workbook(book)
for name,frame in [('Ranking estaciones',summary),('Inventario estaciones',detail),('Metodo estaciones',notes)]:
    if name in wb:
        del wb[name]
    ws=wb.create_sheet(name)
    clean=frame.astype(object).where(pd.notna(frame),None)
    for row in dataframe_to_rows(clean,index=False,header=True):
        ws.append(row)
    ws.freeze_panes='A2'
    ws.auto_filter.ref=ws.dimensions
    for cell in ws[1]:
        cell.font=Font(bold=True,color='FFFFFF')
        cell.fill=PatternFill('solid',fgColor='176B87')
        ws.column_dimensions[cell.column_letter].width=min(45,max(18,len(str(cell.value))+2))
wb.active=wb.sheetnames.index('Ranking estaciones')
for cell in wb['Ranking estaciones'][2]:
    cell.fill=PatternFill('solid',fgColor='D5EEDB')
wb['Metodo estaciones'].column_dimensions['B'].width=115
for row in wb['Metodo estaciones'].iter_rows(min_row=2):
    row[1].alignment=Alignment(wrap_text=True,vertical='top')
    wb['Metodo estaciones'].row_dimensions[row[0].row].height=45
try:
    wb.save(book)
    saved=book
except PermissionError:
    saved=ROOT/'Filtro_8_cuencas_IMERG_con_estaciones.xlsx'
    wb.save(saved)
print('ARCHIVO:',saved)
print(summary[['prioridad','codigo','estacion','meteorologicas_multivariable_total','meteorologicas_multivariable_activas','pluviometricas_pluviograficas_total','pluviometricas_pluviograficas_activas','lluvia_segun_categoria_activas','faltante_max_PQT_1998_2022_pct','pareto_estaciones_faltantes']].to_string(index=False))
